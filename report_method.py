# -*- coding: utf-8 -*-
"""
report_method.py
================
결과물2: 검토용 리포트(xlsx).

목적
----
- 각 값을 '어떤 코드(pandas diff/shift/regression)·어떤 통계기법(피어슨/OLS/차분)·
  어떤 수식'으로 구했는지 추적.
- 신뢰도(R²·p값·잔차진단·VIF·시그마민감도·교차검증)를 한곳에 모음.
- 기법 설명 + 유사 분석 기준 비교.
- '결론 및 유의점'을 자동 생성: 신뢰도 낮음/모순(부호 충돌)/가정 위반을 규칙으로 플래그.

시트 구성
---------
00_방법_요약        : 모델·기법·핵심 신뢰도 지표 요약
01_코드_기법_수식    : 단계별 코드(함수·메서드)·기법·수식·신뢰도확인·유사기법대비
02_신뢰도_진단      : 잔차 정규성/자기상관/등분산 + R² (±Kσ 정당성 근거)
03_시그마민감도      : K별 실제 이상일수 vs 정규기대(%) — K 선택 근거
04_다중공선성_VIF    : 계수 불안정(공선성) 점검
05_상관_차분_비교    : 원시/추세제거 상관 · 차분 동행 · 광고비 차분회귀
06_교차검증          : OLS(±LightGBM) TimeSeriesSplit R²
07_결론_및_유의점    : 자동 플래그(신뢰도 낮음·모순·가정 위반 + 권고)
"""
from __future__ import annotations

import numpy as np
import pandas as pd

from openpyxl.styles import Font, PatternFill, Alignment

import sales_anomaly_engine as E


# ─────────────────────────────────────────────────────────────────────────
# 보조: (라벨, 값) 표에서 안전 추출
# ─────────────────────────────────────────────────────────────────────────
def _get_val(df: pd.DataFrame, label_col: str, label: str, val_col: str):
    """라벨 포함 행의 값(float)을 안전 반환. 없거나 비수치면 None."""
    hit = df[df[label_col].astype(str).str.contains(label, regex=False, na=False)]
    if hit.empty:
        return None
    v = hit.iloc[0][val_col]
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _style_header(ws, n_cols: int, row: int = 1):
    fill = PatternFill("solid", fgColor="2F5530")
    font = Font(color="FFFFFF", bold=True, size=10)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font, cell.alignment = fill, font, align
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


# ─────────────────────────────────────────────────────────────────────────
# 01 코드·기법·수식 매핑(엔진 실제 호출 기준)
# ─────────────────────────────────────────────────────────────────────────
def _code_method_table() -> pd.DataFrame:
    rows = [
        ("일별 매출 집계", "거래단위 → 일단위 Y 생성",
         "df.groupby('결제일자').agg(매출=('실결제금액','sum'))",
         "그룹 합계", "Y = Σ(당일 실결제금액)", "—",
         "피벗테이블 SUMIF와 동일 결과."),
        ("추세항 t", "달력일 기준 추세(행 인덱스 아님)",
         "(df['날짜'] - base).dt.days",
         "선형추세", "t = (날짜 − 기준시작일) [일]",
         "계수 b_t의 p값", "결측일이 있어도 t가 정확(인덱스 방식보다 견고)."),
        ("요일 더미", "요일 효과(월 기준)",
         "pd.get_dummies(요일).drop('요일_월')",
         "범주형 더미", "요일_d ∈ {0,1}, 월요일=기준(drop)",
         "요일 블록 결합 F검정(20b)", "원-핫 인코딩. 기준범주 1개 제거(완전공선성 회피)."),
        ("프로모션 더미", "계획된 이벤트 통제",
         "날짜 ∈ PROMO_SCHEDULE → 1",
         "범주형 더미", "프로모션 ∈ {0,1}",
         "계수 p값", "이상치로 빼지 않고 회귀에 투입(계획된 변동)."),
        ("베이스라인 예측", "추세+요일+프로모션+광고비로 매출 적합",
         "statsmodels OLS(logY ~ X).fit(cov_type='HAC')",
         "다중 선형회귀(OLS) + 로그종속", "ln(Y)=b0+b_t·t+Σγ_d·요일+b_p·프로모+b_a·광고비+ε",
         "R²·각 계수 p값(HAC)", "로그종속=이분산 완화. HAC=Newey-West 자기상관·이분산 강건 표준오차."),
        ("잔차 σ(요일별)", "관리한계 폭 산출",
         "resid.groupby(요일).std(ddof=1)",
         "표본표준편차", "σ_d = sqrt(Σ(e−ē)²/(n−1))",
         "요일별 표본수 N≥10 확인", "요일별 변동성 차이 반영(과검출 교정). N부족 요일은 전체 σ 대체."),
        ("이상치 판정", "예측구간 이탈일 탐지",
         "매출 > 예측·exp(K·σ) or < 예측·exp(−K·σ)",
         "관리한계(±Kσ)", "이상 = |ln Y − ln Ŷ| > K·σ (원단위 Ŷ·exp(±Kσ))",
         "잔차 정규성(JB)·시그마민감도", "SPC 관리도와 동일 개념. K↓=더 많이 잡음."),
        ("잔차 정규성", "±Kσ '몇 %' 해석 정당성",
         "statsmodels jarque_bera(resid)",
         "Jarque-Bera 검정", "JB = n/6·(S² + (K−3)²/4)",
         "p>0.05면 정규 근사", "왜도(S)·첨도(K)로 정규이탈 측정."),
        ("잔차 자기상관", "시계열 의존 점검",
         "durbin_watson(resid) · acorr_ljungbox · resid.shift()",
         "Durbin-Watson / Ljung-Box", "DW=Σ(e_t−e_{t-1})²/Σe_t²  (2≈무상관)",
         "DW 2근처·LB p>0.05", "잔차에 추세/요일 잔존 여부. shift(1)로 lag 비교."),
        ("등분산성", "고매출일 과검출 위험 점검",
         "het_breuschpagan(resid, exog)",
         "Breusch-Pagan 검정", "e² ~ X 회귀의 설명력 검정",
         "p>0.05면 등분산", "이분산이면 σ가 매출수준에 비례 → 로그변환으로 완화."),
        ("수준 상관", "‘매출↑때 광고비↑’ 확인",
         "df['매출'].corr(df['광고비'])",
         "피어슨 상관(레벨)", "r = cov(X,Y)/(sd(X)·sd(Y))",
         "표본수·산점도", "추세 포함. 높은 시점끼리의 동행."),
        ("추세제거 상관", "장기 성장동력 여부",
         "np.polyfit(t,x,1) 잔차끼리 corr",
         "추세제거 후 피어슨", "x̃ = x − (b1·t+b0), r(x̃,ỹ)",
         "—", "+면 성장과 동행, 0/−면 추세교란(공통성장)."),
        ("차분 동행", "단기(늘리면 같이 늚) 동행",
         "df['매출'].diff(); df['광고비'].diff().shift(1)",
         "1차 차분 + (시차) 피어슨", "Δx_t=x_t−x_{t-1}; r(Δx, Δy), r(Δx_{t-1}, Δy_t)",
         "차분 표본수", "추세·레벨 자동 제거. 시차 +면 누적효과 시사."),
        ("광고비 차분회귀", "광고 증액→매출의 기울기·유의",
         "OLS(Δln매출 ~ Δ광고비 당일+전일+요일+프로모).fit(HAC)",
         "차분 회귀 + HAC", "Δln Y = a0 + a1·Δ광고비 + a2·Δ광고비_{t-1} + …",
         "a1·a2 계수 p값", "차분으로 추세·레벨 제거. 계수×100≈광고비1원당 매출%변화."),
        ("VIF", "계수 불안정(공선성) 점검",
         "variance_inflation_factor(exog, i)",
         "분산팽창계수", "VIF_i = 1/(1−R_i²)",
         ">5 주의 / >10 심각", "광고비 음수계수가 공선성 산물인지 판별."),
        ("교차검증", "예측 일반화 성능",
         "TimeSeriesSplit + r2_score",
         "시계열 교차검증", "폴드별 R² 평균(과거→미래 순서 유지)",
         "폴드별 R² 분산", "셔플 금지(시계열 누수 방지). LightGBM 있으면 비교."),
    ]
    cols = ["단계", "목적", "코드(함수·메서드)", "통계기법", "수식",
            "신뢰도 확인", "유사기법 대비/주의"]
    return pd.DataFrame(rows, columns=cols)


# ─────────────────────────────────────────────────────────────────────────
# 07 결론 및 유의점(자동 플래그)
# ─────────────────────────────────────────────────────────────────────────
def _conclusions(result: dict) -> pd.DataFrame:
    res = result["ols"]
    rd = result["resid_diag"]
    flags = []

    def add(item, value, verdict, meaning, advice):
        flags.append({"점검항목": item, "값": value, "판정": verdict,
                      "의미": meaning, "권고": advice})

    # R²
    r2 = float(res.rsquared)
    add("설명력 R²", round(r2, 3),
        "참고" if r2 >= 0.5 else "주의(설명력 낮음)",
        "베이스라인이 매출 변동의 일부만 설명. 단 이상치 판정은 '예측 대비 잔차'라 "
        "R²가 낮아도 잔차진단이 양호하면 ±Kσ 판정 자체는 유효.",
        "예측값 신뢰가 중요하면 변수 보강 검토. 이상탐지 목적이면 02·03 시트로 정당성 확인.")

    # 정규성(JB)
    jb = _get_val(rd, "검정", "Jarque-Bera", "값")
    if jb is not None:
        if jb >= 0.05:
            add("잔차 정규성(JB p)", round(jb, 4), "양호",
                "잔차가 정규에 가까움 → ±Kσ의 '몇 %' 해석이 성립.", "추가 조치 불필요.")
        else:
            add("잔차 정규성(JB p)", round(jb, 4), "주의",
                "잔차가 정규에서 벗어남 → ±Kσ의 정규기대 %는 부정확할 수 있음.",
                "K는 03_시그마민감도의 '실제비율'을 보고 결정(정규기대 맹신 금지).")

    # 자기상관(DW)
    dw = _get_val(rd, "검정", "Durbin-Watson", "값")
    if dw is not None:
        if 1.5 <= dw <= 2.5:
            add("잔차 자기상관(DW)", round(dw, 3), "양호",
                "잔차 자기상관 거의 없음(2 근처).", "추가 조치 불필요.")
        else:
            add("잔차 자기상관(DW)", round(dw, 3), "주의",
                "잔차에 자기상관 존재(시계열 의존) → σ·표준오차 과소평가 가능.",
                f"현재 HAC(maxlags={E.HAC_MAXLAGS}) 적용 중. "
                f"전일매출 항(USE_LAG1) 재도입 검토(현재 {E.USE_LAG1}).")

    # 등분산(BP)
    bp = _get_val(rd, "검정", "Breusch-Pagan", "값")
    if bp is not None:
        if bp >= 0.05:
            add("등분산성(BP p)", round(bp, 4), "양호",
                "등분산(잔차 분산 일정).", "추가 조치 불필요.")
        else:
            add("등분산성(BP p)", round(bp, 4), "주의",
                "이분산 → 고매출일 과검출 위험. 로그종속으로 일부 완화 중.",
                f"로그변환 적용 상태(USE_LOG_TARGET={E.USE_LOG_TARGET}). 잔차 vs 적합값 산점 확인.")

    # VIF
    vif = result.get("vif")
    if vif is not None and "VIF" in vif.columns and len(vif):
        maxv = float(pd.to_numeric(vif["VIF"], errors="coerce").max())
        who = vif.iloc[pd.to_numeric(vif["VIF"], errors="coerce").idxmax()]["변수"]
        verdict = "경고(심각)" if maxv > 10 else ("주의" if maxv > 5 else "양호")
        add(f"다중공선성 VIF(최대: {who})", round(maxv, 2), verdict,
            "VIF가 높은 변수는 계수가 불안정(부호·크기 신뢰 저하).",
            "VIF>10이면 해당 변수 계수 해석 보류. 변수 통합/제거 검토.")

    # 부호 모순: 광고비 회귀계수 vs 원시상관
    coef = result["coef"]
    cr = result.get("correlations")
    ad_coef = _get_val(coef, "변수", "광고비", "계수")
    raw_corr = None
    if cr is not None and "변수" in cr.columns:
        hit = cr[cr["변수"] == "광고비"]
        if not hit.empty and "매출_상관(원시)" in cr.columns:
            try:
                raw_corr = float(hit.iloc[0]["매출_상관(원시)"])
            except (TypeError, ValueError):
                raw_corr = None
    if ad_coef is not None and raw_corr is not None:
        if ad_coef < 0 and raw_corr > 0:
            add("광고비: 회귀계수 vs 원시상관", f"계수 {ad_coef:.2e} / 원시상관 {raw_corr:+.2f}",
                "경고(모순)",
                "통제후 회귀계수는 음수인데 원시상관은 양수 → 추세교란·역인과 가능. "
                "'광고비가 매출을 올린다/내린다'로 단정 불가.",
                "23b_광고비_차분회귀의 부호·유의로 재확인. 확정은 소규모 축소실험 필요.")
        else:
            add("광고비: 회귀계수 vs 원시상관", f"계수 {ad_coef:.2e} / 원시상관 {raw_corr:+.2f}",
                "참고", "회귀계수와 원시상관 부호가 모순되지 않음.",
                "그래도 관측데이터의 인과 단정은 주의(실험으로만 확정).")

    # 광고비 차분회귀 당일
    adr = result.get("ad_diff_reg")
    if adr is not None and "변수" in adr.columns:
        hit = adr[adr["변수"] == "Δ광고비_당일"]
        if not hit.empty:
            try:
                c = float(hit.iloc[0]["계수"]); p = float(hit.iloc[0]["p값"])
                sig = "유의" if p < 0.05 else "비유의"
                add("광고비 차분효과(당일)", f"계수 {c:.2e}, p={p:.3f}",
                    "참고", f"광고비 증액의 당일 매출효과는 {('양수' if c>0 else '음수')}·{sig}.",
                    "비유의면 '단기 매출 견인 근거 약함'. 시차(전일) 계수도 함께 볼 것.")
            except (TypeError, ValueError):
                pass

    # 시그마: 실제비율 vs 정규기대(현재 K)
    ss = result.get("sigma_sens")
    if ss is not None and "현재선택" in ss.columns:
        cur = ss[ss["현재선택"].astype(str).str.contains("현재", na=False)]
        if not cur.empty:
            act = float(cur.iloc[0]["비율(%)"]); exp = float(cur.iloc[0]["정규기대(%)"])
            diff = abs(act - exp)
            verdict = "주의" if diff >= max(3.0, 0.5 * exp) else "양호"
            add(f"이상치 실제비율 vs 정규기대(±{E.SIGMA_K:.2f}σ)",
                f"실제 {act:.1f}% / 기대 {exp:.2f}%", verdict,
                "둘이 비슷하면 ±Kσ 기준이 데이터에 부합. 크게 다르면 정규가정 부적합.",
                "차이가 크면 정규기대% 대신 '실제비율'로 K를 정할 것.")

    out = pd.DataFrame(flags)
    out.attrs["note"] = ("규칙 기반 자동 진단. '경고'는 우선 확인, '주의'는 해석 시 유의, "
                         "'참고/양호'는 통과. 관측데이터 특성상 인과(특히 광고비)는 "
                         "실험 없이 단정하지 않음.")
    return out


# ─────────────────────────────────────────────────────────────────────────
# 메인
# ─────────────────────────────────────────────────────────────────────────
def build_method_report(result: dict, out_path: str) -> str:
    res = result["ols"]
    rd = result["resid_diag"]

    jb = _get_val(rd, "검정", "Jarque-Bera", "값")
    dw = _get_val(rd, "검정", "Durbin-Watson", "값")
    bp = _get_val(rd, "검정", "Breusch-Pagan", "값")
    maxv = None
    if result.get("vif") is not None and "VIF" in result["vif"].columns and len(result["vif"]):
        maxv = float(pd.to_numeric(result["vif"]["VIF"], errors="coerce").max())

    summary = pd.DataFrame([
        ("모델", "로그매출 다중 OLS · 추세 + 요일 + 프로모션 + 광고비"
                 + (" + 전일매출(lag1)" if E.USE_LAG1 else "")),
        ("강건 표준오차", f"Newey-West(HAC), maxlags={E.HAC_MAXLAGS}"),
        ("종속변수 변환", "로그(ln) 적용" if E.USE_LOG_TARGET else "원단위"),
        ("관리한계", f"예측 ±{E.SIGMA_K:.2f}σ"
                    + (" · 요일별 σ" if E.USE_WEEKDAY_SIGMA else " · 전체 σ")),
        ("설명력 R²", round(float(res.rsquared), 4)),
        ("잔차 σ(로그)", round(float(result["sigma"]), 4)),
        ("잔차 정규성(JB p)", round(jb, 4) if jb is not None else "—"),
        ("잔차 자기상관(DW)", round(dw, 3) if dw is not None else "—"),
        ("등분산성(BP p)", round(bp, 4) if bp is not None else "—"),
        ("최대 VIF", round(maxv, 2) if maxv is not None else "—"),
        ("핵심 안내",
         "값별 산출 코드·기법·수식은 01 시트, 신뢰도 근거는 02·03·04·06, "
         "신뢰도 낮음·모순은 07_결론_및_유의점에서 자동 플래그."),
    ], columns=["항목", "값"])

    code_tbl = _code_method_table()
    conclusions = _conclusions(result)

    with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="00_방법_요약", index=False)
        code_tbl.to_excel(xw, sheet_name="01_코드_기법_수식", index=False)
        rd.to_excel(xw, sheet_name="02_신뢰도_진단", index=False)
        result["sigma_sens"].to_excel(xw, sheet_name="03_시그마민감도", index=False)
        result["vif"].to_excel(xw, sheet_name="04_다중공선성_VIF", index=False)

        # 05 상관/차분 비교: 세 표를 한 시트에 위아래로
        ws5 = xw.book.create_sheet("05_상관_차분_비교")
        r = 1
        for title, key in [("[원시/추세제거 상관]", "correlations"),
                           ("[변수효과 비교(수준/차분)]", "var_effect"),
                           ("[광고비 차분회귀]", "ad_diff_reg")]:
            df_ = result.get(key)
            if df_ is None or df_.empty:
                continue
            ws5.cell(row=r, column=1, value=title).font = Font(bold=True, size=11)
            r += 1
            df_.to_excel(xw, sheet_name="05_상관_차분_비교", index=False, startrow=r - 1)
            note = df_.attrs.get("note", "")
            r += len(df_) + 2
            if note:
                ws5.cell(row=r, column=1, value="해설:")
                ws5.cell(row=r, column=2, value=note)
                r += 2

        result["cv"].to_excel(xw, sheet_name="06_교차검증", index=False)
        conclusions.to_excel(xw, sheet_name="07_결론_및_유의점", index=False)

        # 노트 첨부
        for sht, df_ in [("02_신뢰도_진단", rd),
                         ("03_시그마민감도", result["sigma_sens"]),
                         ("04_다중공선성_VIF", result["vif"]),
                         ("07_결론_및_유의점", conclusions)]:
            note = df_.attrs.get("note", "")
            if note:
                ws = xw.sheets[sht]
                ws.cell(row=ws.max_row + 2, column=1, value="해설:")
                ws.cell(row=ws.max_row, column=2, value=note)

        for sht, ncol in [("01_코드_기법_수식", code_tbl.shape[1]),
                          ("02_신뢰도_진단", rd.shape[1]),
                          ("03_시그마민감도", result["sigma_sens"].shape[1]),
                          ("04_다중공선성_VIF", result["vif"].shape[1]),
                          ("06_교차검증", result["cv"].shape[1]),
                          ("07_결론_및_유의점", conclusions.shape[1])]:
            if sht in xw.sheets and ncol > 0:
                _style_header(xw.sheets[sht], ncol)

    print(f"[검토용] 저장 완료: {out_path}")
    return out_path
